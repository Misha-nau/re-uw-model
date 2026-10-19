"""Excel export of an underwriting result (values, not formulas).

For a formula-driven cross-check see backend/scripts/make_crosscheck_xlsx.py.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from re_uw.labels import METRICS
from re_uw.returns import UnderwritingResult

from .schemas import Table, df_to_table

FONT = "Arial"
HEAD = PatternFill("solid", start_color="D9E1F2")
FMT = {"money": '#,##0;(#,##0);"-"', "pct": '0.00%;(0.00%);"-"', "ratio": '0.00"x"',
       "psf": '#,##0.00', "num": '#,##0.00'}


def _write_table(ws, table: Table, title: str, start_row: int) -> int:
    ws.cell(start_row, 1, title).font = Font(name=FONT, bold=True, size=12)
    r = start_row + 1
    ws.cell(r, 1, "").fill = HEAD
    for j, c in enumerate(table.columns, start=2):
        cell = ws.cell(r, j, c)
        cell.font, cell.fill = Font(name=FONT, bold=True), HEAD
    for row in table.rows:
        r += 1
        ws.cell(r, 1, row.label).font = Font(name=FONT, bold=row.emphasis)
        for j, v in enumerate(row.values, start=2):
            cell = ws.cell(r, j, v)
            cell.font = Font(name=FONT, bold=row.emphasis)
            cell.number_format = FMT[row.kind]
    return r + 2


def result_to_xlsx(res: UnderwritingResult) -> bytes:
    inp = res.inputs
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws["A1"] = inp.scenario.name
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = (f"{inp.property.name}. Currency: {inp.analysis.currency}. "
                "Values exported from the model (no formulas).")
    ws["A2"].font = Font(name=FONT, italic=True, color="808080")
    r = 4
    for key, (label, fmt, _group) in METRICS.items():
        v = res.metrics.get(key)
        ws.cell(r, 1, label).font = Font(name=FONT)
        c = ws.cell(r, 2, v)
        c.font, c.number_format = Font(name=FONT), FMT[fmt]
        r += 1

    start = inp.analysis.start_year
    sheets = [
        ("Operating CF", [("Operating cash flow", df_to_table(res.operating.table()))]
         + [(c.name, df_to_table(res.operating.component_table(c.name))) for c in inp.components]),
        ("Returns", [("Cash flows", df_to_table(res.cash_flow_table())),
                     ("Debt schedule", df_to_table(res.debt_table())),
                     ("Equity waterfall", df_to_table(res.waterfall.table(start)))]),
    ]
    for name, tables in sheets:
        sh = wb.create_sheet(name)
        sh.column_dimensions["A"].width = 44
        for j in range(2, 12):
            sh.column_dimensions[get_column_letter(j)].width = 15
        row = 1
        for title, t in tables:
            row = _write_table(sh, t, title, row)

    su = wb.create_sheet("Sources & uses")
    su.column_dimensions["B"].width = 30
    su.column_dimensions["C"].width = 18
    for i, h in enumerate(["Side", "Item", "Amount", "% of total"], start=1):
        c = su.cell(1, i, h)
        c.font, c.fill = Font(name=FONT, bold=True), HEAD
    for i, row in enumerate(res.sources_uses.table().itertuples(index=False), start=2):
        su.cell(i, 1, row.side).font = Font(name=FONT)
        su.cell(i, 2, row.item).font = Font(name=FONT, bold=row.item.startswith("Total"))
        c = su.cell(i, 3, row.amount)
        c.number_format, c.font = FMT["money"], Font(name=FONT)
        c = su.cell(i, 4, row[3])
        c.number_format, c.font = FMT["pct"], Font(name=FONT)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

"""Build an Excel cross-check of the Phase 4 engine for the sample scenario.

Run from backend/:  python scripts/make_crosscheck_xlsx.py [scenario.json] [out.xlsx]

Inputs (blue) are copied from the scenario; NOI and CF before debt come from the Phase 3
engine (already reviewed). Everything from loan sizing to IRR is rebuilt with live Excel
formulas, and compared with the Python result in the "Python" / "Diff" columns.

Scope: interest-only debt, asset-deal stamp duty (price > S$1.5m), forward-NOI exit,
any AM fee basis. Other configurations are covered by the pytest suite.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from re_uw.returns import underwrite
from re_uw.scenarios import PROJECT_ROOT, SCENARIOS_DIR, load_scenario

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
GREY = Font(name=FONT, italic=True, color="808080")
HEAD_FILL = PatternFill("solid", start_color="D9E1F2")
KEY_FILL = PatternFill("solid", start_color="FFFF00")
MONEY = '#,##0;(#,##0);"-"'
PCT = '0.00%;(0.00%);"-"'
MULT = '0.00"x"'


def build(scenario: Path, out: Path) -> Path:
    inp = load_scenario(scenario)
    if inp.debt.amortization_years:
        raise SystemExit("cross-check workbook supports interest-only debt only")
    res = underwrite(inp)
    n = inp.analysis.hold_period_years
    ocf = res.operating
    m = res.metrics

    wb = Workbook()
    ws = wb.active
    ws.title = "Cross-check"
    ws.column_dimensions["A"].width = 46
    for c in range(2, 2 + n + 3):
        ws.column_dimensions[get_column_letter(c)].width = 15

    ws["A1"] = f"Phase 4 cross-check: {inp.scenario.name}"
    ws["A1"].font = TITLE
    ws["A2"] = ("Blue = input copied from the scenario / Phase 3 engine. Black = Excel formula. "
                "Grey = Python result for comparison. Currency: SGD.")
    ws["A2"].font = GREY

    # ---------------- inputs ---------------------------------------------------------
    r = 4
    ws.cell(r, 1, "Inputs").font = BOLD
    ws.cell(r, 1).fill = HEAD_FILL
    a, d, x, f = inp.acquisition, inp.debt, inp.exit, inp.fees
    inputs = [
        ("price", "Purchase price", a.purchase_price, MONEY),
        ("oc", "Other acquisition costs %", a.other_costs_pct, PCT),
        ("capex0", "Upfront capex", a.upfront_capex, MONEY),
        ("ltv", "Target LTV", d.target_ltv if d.enabled else 0.0, PCT),
        ("dscr", "Min DSCR (sizing)", d.min_dscr or 0.0, MULT),
        ("rate", "All-in interest rate (base + margin)", d.all_in_rate(1), PCT),
        ("lfee", "Loan fee %", d.upfront_fee_pct, PCT),
        ("acqfee", "Acquisition fee %", f.acquisition_fee_pct if f.enabled else 0.0, PCT),
        ("amfee", "Asset mgmt fee % (basis below)", f.asset_mgmt_fee_pct if f.enabled else 0.0,
         PCT),
        ("divfee", "Divestment fee %", f.divestment_fee_pct if f.enabled else 0.0, PCT),
        ("cap", "Exit cap rate (forward NOI)", x.exit_cap_rate, PCT),
        ("sell", "Selling costs %", x.selling_cost_pct, PCT),
        ("disc", "Discount rate (NPV)", inp.analysis.discount_rate, PCT),
    ]
    ref: dict[str, str] = {}
    for key, label, value, fmt in inputs:
        r += 1
        ws.cell(r, 1, label).font = BLACK
        c = ws.cell(r, 2, value)
        c.font, c.number_format = BLUE, fmt
        ref[key] = f"$B${r}"
    r += 1
    ws.cell(r, 1, "Asset mgmt fee basis").font = BLACK
    ws.cell(r, 2, f.asset_mgmt_fee_basis).font = BLUE
    basis_cell = f"$B${r}"
    if any(abs(d.all_in_rate(t) - d.all_in_rate(1)) > 1e-12 for t in range(1, n + 1)):
        raise SystemExit("cross-check workbook assumes a flat base rate")

    # ---------------- year grid: NOI / CFBD inputs -----------------------------------
    r += 2
    hdr = r
    ws.cell(hdr, 1, "Year").font = BOLD
    for t in range(n + 2):  # Y0..Y(N+1)
        c = ws.cell(hdr, 2 + t, f"Y{t} {inp.analysis.start_year + t - 1}")
        c.font, c.fill, c.alignment = BOLD, HEAD_FILL, Alignment(horizontal="center")
    ws.cell(hdr, 1).fill = HEAD_FILL

    def col(t: int) -> str:  # year t -> column letter
        return get_column_letter(2 + t)

    rows: dict[str, int] = {}

    def row(key: str, label: str, values=None, formulas=None, fmt=MONEY, font=BLACK):
        nonlocal r
        r += 1
        rows[key] = r
        ws.cell(r, 1, label).font = BOLD if font is BOLD else BLACK
        for t in range(n + 2):
            v = (values[t] if values is not None else None)
            fm = (formulas(t) if formulas is not None else None)
            if v is None and fm is None:
                continue
            c = ws.cell(r, 2 + t, fm if fm is not None else v)
            c.font = (BLUE if values is not None else font)
            c.number_format = fmt

    noi_vals = [None] + [float(v) for v in ocf.noi]
    cfbd_vals = [None] + [float(v) for v in ocf.cf_before_debt[:n]] + [None]
    row("noi", "NOI (Phase 3 engine)", values=noi_vals)
    row("cfbd", "CF before debt (Phase 3 engine)", values=cfbd_vals)

    # ---------------- acquisition & financing (formulas) -----------------------------
    r += 1
    calc = {}

    def scalar(key: str, label: str, formula: str, py: float | None, fmt=MONEY, key_fill=False):
        nonlocal r
        r += 1
        ws.cell(r, 1, label).font = BLACK
        c = ws.cell(r, 2, formula)
        c.font, c.number_format = BLACK, fmt
        if key_fill:
            c.fill = KEY_FILL
        if py is not None:
            p = ws.cell(r, 3, py)
            p.font, p.number_format = GREY, fmt
            dcell = ws.cell(r, 4, f"=B{r}-C{r}")
            dcell.font, dcell.number_format = BLACK, ('0.000000%' if fmt == PCT else '0.00')
        calc[key] = f"$B${r}"

    r += 1
    ws.cell(r, 1, "Acquisition, financing, exit").font = BOLD
    ws.cell(r, 1).fill = HEAD_FILL
    for cc, lab in ((2, "Excel"), (3, "Python"), (4, "Diff")):
        ws.cell(r, cc, lab).font = BOLD
        ws.cell(r, cc).fill = HEAD_FILL

    su, ex = res.sources_uses, res.exit
    noi1 = f"{col(1)}{rows['noi']}"
    noi_fwd = f"{col(n + 1)}{rows['noi']}"
    scalar("sd", "Stamp duty (BSD: 5% x price - 30,400)", f"={ref['price']}*5%-30400",
           su.stamp_duty)
    scalar("ocost", "Other acquisition costs", f"={ref['price']}*{ref['oc']}", su.other_costs)
    scalar("ltvamt", "Loan at target LTV", f"={ref['ltv']}*{ref['price']}", res.sizing.ltv_amount)
    scalar("dscramt", "Loan at min DSCR (IO: NOI1 / (DSCR x rate))",
           f"=IF({ref['dscr']}>0,{noi1}/({ref['dscr']}*{ref['rate']}),{calc['ltvamt']})",
           res.sizing.dscr_amount)
    scalar("loan", "Loan (lower of the two)", f"=MIN({calc['ltvamt']},{calc['dscramt']})",
           su.loan, key_fill=True)
    scalar("lfee_amt", "Loan fee", f"={calc['loan']}*{ref['lfee']}", su.loan_fee)
    scalar("acq_amt", "Acquisition fee", f"={ref['price']}*{ref['acqfee']}", su.acquisition_fee)
    scalar("propcost", "Property cost (price + duty + costs + capex)",
           f"={ref['price']}+{calc['sd']}+{calc['ocost']}+{ref['capex0']}", su.property_cost)
    scalar("uses", "Total uses", f"={calc['propcost']}+{calc['lfee_amt']}+{calc['acq_amt']}",
           su.total_uses)
    scalar("equity", "Equity at closing", f"={calc['uses']}-{calc['loan']}", su.equity,
           key_fill=True)
    scalar("gross", "Exit value (forward NOI / exit cap)", f"={noi_fwd}/{ref['cap']}",
           ex.gross_value)
    scalar("sellc", "Selling costs", f"={calc['gross']}*{ref['sell']}", ex.selling_costs)
    scalar("divf", "Divestment fee", f"={calc['gross']}*{ref['divfee']}", ex.divestment_fee)

    # ---------------- annual cash flows (formulas) -----------------------------------
    r += 1
    yrs = range(0, n + 1)

    def frow(key: str, label: str, fn, bold=False):
        nonlocal r
        r += 1
        rows[key] = r
        ws.cell(r, 1, label).font = BOLD if bold else BLACK
        for t in yrs:
            c = ws.cell(r, 2 + t, fn(t))
            c.font, c.number_format = (BOLD if bold else BLACK), MONEY

    def at(key: str, t: int) -> str:
        return f"{col(t)}{rows[key]}"

    frow("interest", "Interest (IO)", lambda t: 0 if t == 0 else f"={calc['loan']}*{ref['rate']}")
    frow("am", "Asset mgmt fee",
         lambda t: 0 if t == 0 else (
             f'=IF({basis_cell}="gav",{ref["amfee"]}*{col(t + 1)}{rows["noi"]}/{ref["cap"]},'
             f'IF({basis_cell}="purchase_price",{ref["amfee"]}*{ref["price"]},'
             f'{ref["amfee"]}*{calc["equity"]}))'))
    frow("unlev", "Unlevered cash flow",
         lambda t: f"=-{calc['propcost']}" if t == 0 else
         (f"={at('cfbd', t)}+{calc['gross']}-{calc['sellc']}" if t == n else f"={at('cfbd', t)}"),
         bold=True)
    frow("levg", "Levered cash flow (gross of fund fees)",
         lambda t: f"=-({calc['equity']}-{calc['acq_amt']})" if t == 0 else
         (f"={at('unlev', t)}-{at('interest', t)}-{calc['loan']}" if t == n
          else f"={at('unlev', t)}-{at('interest', t)}"), bold=True)
    frow("levn", "Equity cash flow (net of fees)",
         lambda t: f"=-{calc['equity']}" if t == 0 else
         (f"={at('levg', t)}-{at('am', t)}-{calc['divf']}" if t == n
          else f"={at('levg', t)}-{at('am', t)}"), bold=True)

    # compare cash flow rows with Python
    r += 1
    py_rows = [("unlev", res.unlevered_cf), ("levg", res.levered_gross_cf),
               ("levn", res.levered_net_cf)]
    for key, arr in py_rows:
        r += 1
        ws.cell(r, 1, f"  Python: {ws.cell(rows[key], 1).value}").font = GREY
        for t in yrs:
            c = ws.cell(r, 2 + t, float(arr[t]))
            c.font, c.number_format = GREY, MONEY
        rows[key + "_py"] = r

    # ---------------- returns ------------------------------------------------------------
    r += 1
    r += 1
    ws.cell(r, 1, "Returns").font = BOLD
    ws.cell(r, 1).fill = HEAD_FILL
    for cc, lab in ((2, "Excel"), (3, "Python"), (4, "Diff")):
        ws.cell(r, cc, lab).font = BOLD
        ws.cell(r, cc).fill = HEAD_FILL

    def rng(key: str) -> str:
        return f"{col(0)}{rows[key]}:{col(n)}{rows[key]}"

    def npv_f(key: str) -> str:
        return f"={col(0)}{rows[key]}+NPV({ref['disc']},{col(1)}{rows[key]}:{col(n)}{rows[key]})"

    def mult_f(key: str) -> str:
        return f'=SUMIF({rng(key)},">0")/-SUMIF({rng(key)},"<0")'

    scalar("uirr", "Unlevered IRR", f"=IRR({rng('unlev')})", m["unlevered_irr"], PCT, True)
    scalar("umult", "Unlevered multiple", mult_f("unlev"), m["unlevered_multiple"], MULT)
    scalar("unpv", "Unlevered NPV", npv_f("unlev"), m["unlevered_npv"])
    scalar("lgirr", "Levered IRR (gross of fund fees)", f"=IRR({rng('levg')})",
           m["levered_irr_gross"], PCT, True)
    scalar("lirr", "Equity IRR (net of fees)", f"=IRR({rng('levn')})", m["levered_irr"], PCT, True)
    scalar("lmult", "Equity multiple (net)", mult_f("levn"), m["levered_multiple"], MULT)
    scalar("lnpv", "Equity NPV (net)", npv_f("levn"), m["levered_npv"])
    scalar("dscr1", "Year 1 DSCR (NOI1 / interest)", f"={noi1}/{at('interest', 1)}",
           m["dscr_year1"], MULT)
    scalar("entry", "Entry cap rate (NOI1 / price)", f"={noi1}/{ref['price']}",
           m["entry_cap_rate"], PCT)

    r += 2
    ws.cell(r, 1, "Largest absolute cash flow difference (Excel - Python)").font = BOLD
    diffs = []
    for k, _ in py_rows:
        for t in yrs:
            diffs.append(f"ABS({col(t)}{rows[k]}-{col(t)}{rows[k + '_py']})")
    c = ws.cell(r, 2, "=MAX(" + ",".join(diffs) + ")")
    c.font, c.number_format, c.fill = BOLD, "0.000000", KEY_FILL
    ws.cell(r, 3, "should be ~0 (rounding only)").font = GREY

    ws.freeze_panes = ws.cell(hdr + 1, 2)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    scen = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        SCENARIOS_DIR / "sample_sg_cbd_office_retail.json")
    dest = Path(sys.argv[2]) if len(sys.argv) > 2 else (
        PROJECT_ROOT / "reports" / "phase4_crosscheck.xlsx")
    print(f"wrote {build(scen, dest)}")
